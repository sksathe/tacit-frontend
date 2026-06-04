from __future__ import annotations

from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def write_logistics_workbook(*, payload: Dict[str, Any], output_xlsx_path: str) -> None:
  wb = Workbook()
  ws0 = wb.active
  ws0.title = "Summary"
  row = 1
  ws0.cell(row=row, column=1, value="Document type")
  ws0.cell(row=row, column=2, value=str(payload.get("documentType") or ""))
  row += 1
  if payload.get("documentTypeLabel"):
    ws0.cell(row=row, column=1, value="Label")
    ws0.cell(row=row, column=2, value=str(payload.get("documentTypeLabel")))
    row += 1
  if payload.get("confidence") is not None:
    ws0.cell(row=row, column=1, value="Confidence")
    ws0.cell(row=row, column=2, value=payload.get("confidence"))
    row += 1

  summary = payload.get("summary") or {}
  if summary.get("title"):
    row += 1
    ws0.cell(row=row, column=1, value="Summary title")
    ws0.cell(row=row, column=2, value=str(summary.get("title")))
    row += 1
  bullets: List[str] = summary.get("bullets") or []
  if bullets:
    ws0.cell(row=row, column=1, value="Summary bullets")
    row += 1
    for b in bullets:
      ws0.cell(row=row, column=1, value=str(b))
      row += 1

  row += 1
  ws0.cell(row=row, column=1, value="Field groups")
  row += 1
  for g in payload.get("fieldGroups") or []:
    if not isinstance(g, dict):
      continue
    ws0.cell(row=row, column=1, value=str(g.get("label") or g.get("id")))
    row += 1
    for f in g.get("fields") or []:
      if not isinstance(f, dict):
        continue
      ws0.cell(row=row, column=1, value=str(f.get("label")))
      ws0.cell(row=row, column=2, value=str(f.get("value") or ""))
      row += 1

  ws1 = wb.create_sheet("Tables")
  r = 1
  for tbl in payload.get("tables") or []:
    if not isinstance(tbl, dict):
      continue
    ws1.cell(row=r, column=1, value=str(tbl.get("title") or tbl.get("id")))
    r += 1
    cols: List[str] = list(tbl.get("columns") or [])
    if cols:
      for i, c in enumerate(cols):
        ws1.cell(row=r, column=i + 1, value=str(c))
      r += 1
    for row_obj in tbl.get("rows") or []:
      if not isinstance(row_obj, dict):
        continue
      for i, c in enumerate(cols):
        ws1.cell(row=r, column=i + 1, value=str(row_obj.get(c, "")))
      r += 1
    r += 1

  ws2 = wb.create_sheet("Warnings")
  ws2.cell(row=1, column=1, value="Severity")
  ws2.cell(row=1, column=2, value="Code")
  ws2.cell(row=1, column=3, value="Message")
  wr = 2
  for w in payload.get("warnings") or []:
    if not isinstance(w, dict):
      continue
    ws2.cell(row=wr, column=1, value=str(w.get("severity")))
    ws2.cell(row=wr, column=2, value=str(w.get("code") or ""))
    ws2.cell(row=wr, column=3, value=str(w.get("message")))
    wr += 1

  for ws in (ws0, ws1, ws2):
    for col in range(1, 12):
      ws.column_dimensions[get_column_letter(col)].width = 18

  wb.save(output_xlsx_path)
